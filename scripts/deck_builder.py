"""
deck_builder.py - Master Slide Builder Primitives for AWS Well-Architected Presentation
Implements an Apple/AWS executive sober aesthetic:
- Off-white canvas background (#F8F9FA)
- Charcoal black ink (#0F172A)
- Slate secondary text (#475569)
- Soft hairline borders (#E2E8F0)
- Purposeful semantic accents:
  - AWS Amber (#D97706)
  - Enterprise Navy (#1E293B)
  - Resilient Emerald (#059669 / #10B981)
  - Fragile Crimson (#DC2626 / #EF4444)
  - FinOps Purple (#7C3AED)
  - Cobalt Blue (#2563EB)
  - Sustainability Teal (#0D9488)
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 16:9 Widescreen dimensions
SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5
MARGIN_X = 0.8
CONTENT_WIDTH = SLIDE_WIDTH_IN - (MARGIN_X * 2) # 11.733 in

# Color Palette
HEX_BG = "F8F9FA"          # Off-white warm canvas
HEX_SURFACE = "FFFFFF"     # Pure white card surface
HEX_SURFACE_MUTED = "F1F5F9" # Light slate muted card
HEX_INK = "0F172A"          # Deep charcoal
HEX_SECONDARY = "475569"    # Slate secondary
HEX_MUTED = "64748B"        # Lighter slate
HEX_BORDER = "E2E8F0"       # Subtle card border
HEX_BORDER_DARK = "CBD5E1"  # Slightly darker border
HEX_AMBER = "D97706"        # AWS Amber
HEX_NAVY = "1E293B"         # AWS Dark Navy
HEX_BLUE = "2563EB"         # Enterprise Cobalt Blue
HEX_SUCCESS = "059669"      # Resilient Emerald Green
HEX_SUCCESS_LIGHT = "F0FDF4"
HEX_SUCCESS_BORDER = "86EFAC"
HEX_DANGER = "DC2626"       # Fragile Crimson Red
HEX_DANGER_LIGHT = "FEF2F2"
HEX_DANGER_BORDER = "FCA5A5"
HEX_PURPLE = "7C3AED"       # FinOps Purple
HEX_PURPLE_LIGHT = "F5F3FF"
HEX_PURPLE_BORDER = "C4B5FD"
HEX_TEAL = "0D9488"         # Sustainability Teal
HEX_TEAL_LIGHT = "F0FDFA"
HEX_TEAL_BORDER = "99F6E4"

def hex_to_rgb(hex_str: str) -> RGBColor:
    h = hex_str.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

class MasterDeckBuilder:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_WIDTH_IN)
        self.prs.slide_height = Inches(SLIDE_HEIGHT_IN)
        self.blank_layout = self.prs.slide_layouts[6]
        self.total_slides_count = 140

    def set_total_slides(self, count: int):
        self.total_slides_count = count

    def create_blank_slide(self):
        slide = self.prs.slides.add_slide(self.blank_layout)
        # Background fill
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_WIDTH_IN), Inches(SLIDE_HEIGHT_IN))
        bg.fill.solid()
        bg.fill.fore_color.rgb = hex_to_rgb(HEX_BG)
        bg.line.fill.background()
        return slide

    def add_header(self, slide, eyebrow: str, title: str, banner_text: str = None, banner_style: str = "default"):
        """Adds standard header with eyebrow, title, and optional banner strip with adaptive font sizing."""
        # Eyebrow
        eyebrow_box = slide.shapes.add_textbox(Inches(MARGIN_X), Inches(0.42), Inches(CONTENT_WIDTH), Inches(0.26))
        tf = eyebrow_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = eyebrow.upper()
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb(HEX_AMBER if "MODULE" in eyebrow or "CHAOS" in eyebrow else HEX_BLUE)

        # Adaptive Title Sizing & Height
        t_len = len(title)
        if t_len <= 52:
            title_size = Pt(20)
            title_h = 0.44
        elif t_len <= 80:
            title_size = Pt(16.5)
            title_h = 0.50
        elif t_len <= 110:
            title_size = Pt(15.0)
            title_h = 0.54
        else: # Very long titles (> 110 chars)
            title_size = Pt(13.5)
            title_h = 0.56

        title_box = slide.shapes.add_textbox(Inches(MARGIN_X), Inches(0.70), Inches(CONTENT_WIDTH), Inches(title_h))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
        p2 = tf2.paragraphs[0]
        p2.text = title
        p2.font.name = "Arial"
        p2.font.size = title_size
        p2.font.bold = True
        p2.font.color.rgb = hex_to_rgb(HEX_INK)

        # Optional Banner with Adaptive Positioning
        if banner_text:
            bg_color = HEX_SURFACE_MUTED
            border_color = HEX_BORDER
            text_color = HEX_INK
            prefix = "📌 "
            if banner_style == "danger":
                bg_color = HEX_DANGER_LIGHT
                border_color = HEX_DANGER_BORDER
                text_color = HEX_DANGER
                prefix = "⚠️ "
            elif banner_style == "success":
                bg_color = HEX_SUCCESS_LIGHT
                border_color = HEX_SUCCESS_BORDER
                text_color = HEX_SUCCESS
                prefix = "✅ "
            elif banner_style == "accent":
                bg_color = "EFF6FF"
                border_color = "BFDBFE"
                text_color = HEX_BLUE
                prefix = "💡 "
            elif banner_style == "purple":
                bg_color = HEX_PURPLE_LIGHT
                border_color = HEX_PURPLE_BORDER
                text_color = HEX_PURPLE
                prefix = "📊 "

            # Calculate banner Y position dynamically from title height
            banner_y = max(1.26, 0.70 + title_h + 0.08)

            b_len = len(banner_text)
            if b_len <= 85:
                banner_size = Pt(10.0)
                banner_h = 0.36
            elif b_len <= 140:
                banner_size = Pt(9.0)
                banner_h = 0.40
            else:
                banner_size = Pt(8.5)
                banner_h = 0.44

            # Ensure banner leaves at least 0.12in breathing room before cards start at y=1.95
            if banner_y + banner_h > 1.83:
                banner_y = 1.83 - banner_h

            banner_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(MARGIN_X), Inches(banner_y), Inches(CONTENT_WIDTH), Inches(banner_h)
            )
            banner_shape.fill.solid()
            banner_shape.fill.fore_color.rgb = hex_to_rgb(bg_color)
            banner_shape.line.color.rgb = hex_to_rgb(border_color)
            banner_shape.line.width = Pt(1)

            tf3 = banner_shape.text_frame
            tf3.word_wrap = True
            tf3.margin_left = Inches(0.18)
            tf3.margin_top = Inches(0.04)
            tf3.margin_right = Inches(0.18)
            tf3.margin_bottom = Inches(0.04)
            p3 = tf3.paragraphs[0]
            p3.text = prefix + banner_text
            p3.font.name = "Arial"
            p3.font.size = banner_size
            p3.font.bold = True
            p3.font.color.rgb = hex_to_rgb(text_color)

    def add_footer(self, slide, slide_num: int):
        """Adds consistent, elegant bottom hairline rule and metadata."""
        # Hairline separator
        rule = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(MARGIN_X), Inches(7.04), Inches(CONTENT_WIDTH), Inches(0.015)
        )
        rule.fill.solid()
        rule.fill.fore_color.rgb = hex_to_rgb(HEX_BORDER)
        rule.line.fill.background()

        # Left label
        left_box = slide.shapes.add_textbox(Inches(MARGIN_X), Inches(7.08), Inches(4.5), Inches(0.3))
        tf = left_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "AWS Well-Architected Framework · Enterprise Case Study"
        p.font.name = "Arial"
        p.font.size = Pt(8.5)
        p.font.color.rgb = hex_to_rgb(HEX_MUTED)

        # Center link
        center_box = slide.shapes.add_textbox(Inches(5.3), Inches(7.08), Inches(3.0), Inches(0.3))
        tf_c = center_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        p_c = tf_c.paragraphs[0]
        p_c.alignment = PP_ALIGN.CENTER
        p_c.text = "wafpresentation.vercel.app"
        p_c.font.name = "Arial"
        p_c.font.size = Pt(8.5)
        p_c.font.color.rgb = hex_to_rgb(HEX_MUTED)

        # Right authors & slide counter
        right_box = slide.shapes.add_textbox(Inches(8.3), Inches(7.08), Inches(4.233), Inches(0.3))
        tf_r = right_box.text_frame
        tf_r.word_wrap = True
        tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
        p_r = tf_r.paragraphs[0]
        p_r.alignment = PP_ALIGN.RIGHT
        p_r.text = f"Devarsh Patel (24BCT0267) & Aman Kumar (24BCE0839) · Slide {slide_num} of {self.total_slides_count}"
        p_r.font.name = "Arial"
        p_r.font.size = Pt(8.5)
        p_r.font.bold = True
        p_r.font.color.rgb = hex_to_rgb(HEX_SECONDARY)

    def add_notes(self, slide, notes_text: str):
        """Adds speaker notes for prompter and defense readiness."""
        if not notes_text:
            return
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    def add_card(self, slide, x: float, y: float, w: float, h: float, bg_color: str = HEX_SURFACE, border_color: str = HEX_BORDER, border_width: float = 1.0):
        """Creates a styled card container."""
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = hex_to_rgb(bg_color)
        card.line.color.rgb = hex_to_rgb(border_color)
        card.line.width = Pt(border_width)
        return card

    def add_badge(self, slide, x: float, y: float, w: float, h: float, text: str, bg_color: str, text_color: str, border_color: str = None, font_size: float = 8.5):
        """Adds a small tag/badge."""
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        badge.fill.solid()
        badge.fill.fore_color.rgb = hex_to_rgb(bg_color)
        if border_color:
            badge.line.color.rgb = hex_to_rgb(border_color)
            badge.line.width = Pt(0.75)
        else:
            badge.line.fill.background()
        tf = badge.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = text
        p.font.name = "Arial"
        p.font.size = Pt(font_size)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb(text_color)
        return badge

    def add_table(self, slide, x: float, y: float, w: float, h: float, headers: list, rows: list, col_widths: list = None, font_size: float = 9.5, header_bg: str = HEX_SURFACE_MUTED, header_fg: str = HEX_NAVY):
        """Creates a clean, styled table with borders and alternating rows."""
        row_count = len(rows) + 1
        col_count = len(headers)
        table_shape = slide.shapes.add_table(row_count, col_count, Inches(x), Inches(y), Inches(w), Inches(h))
        tbl = table_shape.table

        # Set column widths
        if col_widths and len(col_widths) == col_count:
            for i, width in enumerate(col_widths):
                tbl.columns[i].width = Inches(width)

        # Style Header
        for col_idx, header_text in enumerate(headers):
            cell = tbl.cell(0, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = hex_to_rgb(header_bg)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.1)
            tf.margin_right = Inches(0.1)
            tf.margin_top = Inches(0.06)
            tf.margin_bottom = Inches(0.06)
            p = tf.paragraphs[0]
            p.text = str(header_text).upper()
            p.font.name = "Arial"
            p.font.size = Pt(font_size)
            p.font.bold = True
            p.font.color.rgb = hex_to_rgb(header_fg)

        # Style Body Rows
        for row_idx, row_data in enumerate(rows):
            is_even = (row_idx % 2 == 0)
            row_bg = HEX_SURFACE if is_even else HEX_SURFACE_MUTED
            for col_idx, val in enumerate(row_data):
                cell = tbl.cell(row_idx + 1, col_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = hex_to_rgb(row_bg)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                tf.margin_left = Inches(0.1)
                tf.margin_right = Inches(0.1)
                tf.margin_top = Inches(0.05)
                tf.margin_bottom = Inches(0.05)
                p = tf.paragraphs[0]
                p.text = str(val)
                p.font.name = "Arial"
                p.font.size = Pt(font_size - 0.5)
                p.font.bold = (col_idx == 0)
                p.font.color.rgb = hex_to_rgb(HEX_INK if col_idx == 0 else HEX_SECONDARY)

        return table_shape
