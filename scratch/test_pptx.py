import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6] # blank layout
slide = prs.slides.add_slide(blank_slide_layout)

# Background
bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(0xF8, 0xF9, 0xFA)
bg.line.color.rgb = RGBColor(0xF8, 0xF9, 0xFA)

# Title
tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.2))
tf = tx_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "AWS Well-Architected Framework: Naive vs Resilient"
p.font.name = "Arial"
p.font.size = Pt(28)
p.font.bold = True
p.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

prs.save("scratch/test_out.pptx")
print("Saved test pptx successfully!")
